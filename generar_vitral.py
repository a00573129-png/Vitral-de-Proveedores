"""
generar_vitral.py — Generador automático del Vitral de Proveedores CLRAH
────────────────────────────────────────────────────────────────────────
Uso:
    python generar_vitral.py

El script lee 'Base_de_datos_Vitral_de_Proveedores.xlsx' y regenera 
'Vitral_Proveedores_CLRAH.html' manteniendo la estructura gráfica moderna.
"""

import json, sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Instalando openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

import warnings
warnings.filterwarnings("ignore")

# ── Configuración ────────────────────────────────────────────
EXCEL_FILE = Path(__file__).parent / "Base_de_datos_Vitral_de_Proveedores.xlsx"
HTML_FILE  = Path(__file__).parent / "Vitral_Proveedores_CLRAH.html"

COL = {
    "nombre":    1,
    "cat":       2,
    "region":    4,
    "trans":     5,
    "email":     6,
    "tel":       7,
    "cert":      8,
    "cap":       9,
    "estado":   10,
    "notas":    12,
    "lat":      13,
    "lng":      14,
    "lugar":    15,
    "prioridad":16,
}

def cat_key(categoria):
    """Mapea la categoría del Excel a las claves estandarizadas del Vitral."""
    if not categoria:
        return "logistica"
    cat = str(categoria).lower()
    
    if "wash" in cat or "agua" in cat or "saneamiento" in cat:
        return "wash"
    if "shelter" in cat or "refugio" in cat:
        return "shelter"
    if "camp" in cat or "campamento" in cat:
        return "camp"
    if "protection" in cat or "protecci" in cat:
        return "protection"
    if "health" in cat or "salud" in cat or "farmac" in cat or "médic" in cat or "medic" in cat:
        return "health"
    if "transporte" in cat or "logística" in cat or "logistica" in cat:
        return "logistica"
    return "otro"

def leer_excel():
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Proveedores"]
    proveedores = []
    
    # Lectura desde la fila 3 (evitamos títulos)
    for r in range(3, ws.max_row + 1):
        nombre = ws.cell(r, COL["nombre"]).value
        if not nombre or str(nombre).strip() == "Nombre del Proveedor":
            continue
            
        lat_raw  = ws.cell(r, COL["lat"]).value
        lng_raw  = ws.cell(r, COL["lng"]).value
        try:
            lat = float(lat_raw)
            lng = float(lng_raw)
        except (TypeError, ValueError):
            print(f"  ⚠ Fila {r} ({nombre}): Lat/Lng sin definir, se coloca en Colón por defecto.")
            lat, lng = 9.3548, -79.8994

        prio_raw = str(ws.cell(r, COL["prioridad"]).value or "").strip().upper()
        prioridad = prio_raw in ("SI", "SÍ", "YES", "TRUE", "1")

        p = {
            "nombre":    str(nombre).strip(),
            "cat":       str(ws.cell(r, COL["cat"]).value or "").strip(),
            "catKey":    cat_key(ws.cell(r, COL["cat"]).value),
            "region":    str(ws.cell(r, COL["region"]).value or "").strip(),
            "trans":     str(ws.cell(r, COL["trans"]).value or "").strip(),
            "email":     str(ws.cell(r, COL["email"]).value or "").strip(),
            "tel":       str(ws.cell(r, COL["tel"]).value or "").strip(),
            "cert":      str(ws.cell(r, COL["cert"]).value or "").strip(),
            "cap":       str(ws.cell(r, COL["cap"]).value or "").strip(),
            "estado":    str(ws.cell(r, COL["estado"]).value or "Potencial").strip(),
            "notas":     str(ws.cell(r, COL["notas"]).value or "").strip(),
            "lat":       lat,
            "lng":       lng,
            "lugar":     str(ws.cell(r, COL["lugar"]).value or "Panamá").strip(),
            "prioridad": prioridad,
        }
        proveedores.append(p)

    print(f"  {len(proveedores)} proveedores leídos del Excel")
    return proveedores

def generar_html(proveedores):
    datos_js = json.dumps(proveedores, ensure_ascii=False, indent=2)
    n = len(proveedores)

    html = f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Vitral de Proveedores — CLRAH</title>

<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Cinzel:wght@600;700;800;900&family=Roboto:ital,wght@0,400;0,500;0,700;1,400&display=swap" rel="stylesheet">

<link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/leaflet/1.9.4/leaflet.min.css"/>
<script src="https://cdnjs.cloudflare.com/ajax/libs/leaflet/1.9.4/leaflet.min.js"></script>

<style>
*{{box-sizing:border-box;margin:0;padding:0;}}

:root{{
  --clrah-red: #E30613;
  --clrah-blue: #0055A5;
  --clrah-gray: #58595B;
  --clrah-bg: #F8FAFC;
  
  --ink: #1E293B;
  --muted: #64748B;
  --rule: #CBD5E1;
  --white: #FFFFFF;

  --col-wash: #0284C7;
  --col-shelter: #16A34A;
  --col-camp: #CA8A04;
  --col-protection: #DC2626;
  --col-health: #7C3AED;
  --col-logistica: #0055A5;
  --col-otro: #64748B;
}}

html, body {{
  height: 100%;
  width: 100%;
  margin: 0;
  padding: 0;
  font-family: 'Roboto', sans-serif;
  background: var(--clrah-bg);
  color: var(--ink);
  display: flex;
  flex-direction: column;
  overflow: hidden;
}}

header {{
  background: var(--white);
  height: 60px;
  padding: 0 24px;
  display: flex;
  align-items: center;
  justify-content: space-between;
  flex-shrink: 0;
  border-bottom: 4px solid var(--clrah-red);
  box-shadow: 0 2px 8px rgba(0,0,0,0.08);
  z-index: 1000;
}}
.h-izq {{ display: flex; align-items: center; gap: 14px; }}
.h-marca {{
  font-family: 'Cinzel', serif;
  font-size: 1.4rem;
  font-weight: 800;
  color: var(--clrah-gray);
  letter-spacing: 0.05em;
  padding-right: 14px;
  border-right: 2px solid var(--rule);
}}
.h-marca span {{ color: var(--clrah-red); }}
.h-titulo {{
  font-family: 'Roboto', sans-serif;
  font-size: 0.85rem;
  font-weight: 700;
  color: var(--clrah-blue);
  text-transform: uppercase;
  letter-spacing: 0.03em;
}}
.h-der {{ display: flex; align-items: center; gap: 10px; }}
.h-badge {{
  background: var(--clrah-blue);
  color: var(--white);
  font-size: 0.72rem;
  font-weight: 700;
  padding: 5px 12px;
  border-radius: 3px;
}}
.h-btn {{
  font-family: 'Roboto', sans-serif;
  font-size: 0.73rem;
  font-weight: 700;
  background: var(--clrah-red);
  color: var(--white);
  border: none;
  padding: 7px 14px;
  border-radius: 3px;
  cursor: pointer;
  transition: background 0.15s;
}}
.h-btn:hover {{ background: #B8040F; }}

.filtros {{
  background: #F1F5F9;
  border-bottom: 1px solid var(--rule);
  padding: 0 24px;
  height: 48px;
  display: flex;
  align-items: center;
  gap: 14px;
  flex-shrink: 0;
}}
.flabel {{
  font-size: 0.65rem;
  font-weight: 700;
  color: var(--clrah-gray);
  text-transform: uppercase;
  letter-spacing: 0.05em;
}}
.fsep {{ width: 1px; height: 20px; background: var(--rule); }}
select, input[type=text] {{
  font-family: 'Roboto', sans-serif;
  font-size: 0.78rem;
  color: var(--ink);
  background: var(--white);
  border: 1px solid var(--rule);
  border-radius: 3px;
  padding: 5px 10px;
  outline: none;
}}
select:focus, input:focus {{ border-color: var(--clrah-blue); }}
input[type=text] {{ width: 190px; }}
#rcount {{ margin-left: auto; font-size: 0.75rem; font-weight: 600; color: var(--muted); }}
#rcount span {{ color: var(--clrah-blue); font-weight: 700; }}

.grid {{
  display: grid;
  grid-template-columns: 360px 1fr;
  flex: 1;
  height: calc(100vh - 108px);
  width: 100%;
  overflow: hidden;
}}

.sidebar {{
  background: var(--white);
  border-right: 1px solid var(--rule);
  overflow-y: auto;
  height: 100%;
}}
.item {{
  padding: 14px 18px;
  border-bottom: 1px solid var(--rule);
  cursor: pointer;
  transition: background 0.15s;
  position: relative;
}}
.item:hover {{ background: #F8FAFC; }}
.item.activo {{
  background: #E0F2FE;
  border-left: 5px solid var(--clrah-blue);
  padding-left: 13px;
}}
.iname {{
  font-family: 'Cinzel', serif;
  font-size: 0.92rem;
  font-weight: 700;
  color: #0F172A;
  margin-bottom: 6px;
  line-height: 1.3;
  padding-right: 60px;
}}
.imeta {{ display: flex; align-items: center; gap: 8px; flex-wrap: wrap; }}
.cpill {{
  font-size: 0.62rem;
  font-weight: 700;
  padding: 2px 8px;
  border-radius: 3px;
  color: #FFFFFF;
}}

.cpill.wash {{ background: var(--col-wash); }}
.cpill.shelter {{ background: var(--col-shelter); }}
.cpill.camp {{ background: var(--col-camp); }}
.cpill.protection {{ background: var(--col-protection); }}
.cpill.health {{ background: var(--col-health); }}
.cpill.logistica {{ background: var(--col-logistica); }}
.cpill.otro {{ background: var(--col-otro); }}

.ilugar {{ font-size: 0.72rem; color: var(--muted); }}
.pmark {{
  position: absolute;
  top: 14px;
  right: 14px;
  font-size: 0.6rem;
  font-weight: 800;
  color: var(--clrah-red);
  text-transform: uppercase;
}}

#mapa {{
  width: 100%;
  height: 100%;
  min-height: 100%;
  background: #E2E8F0;
}}

.lmapa {{
  background: var(--white);
  border-left: 4px solid var(--clrah-blue);
  border-radius: 3px;
  padding: 12px 14px;
  font-size: 0.72rem;
  line-height: 1.8;
  box-shadow: 0 4px 12px rgba(0,0,0,0.15);
}}
.lmapa b {{
  font-family: 'Cinzel', serif;
  display: block;
  font-size: 0.75rem;
  font-weight: 700;
  color: var(--clrah-gray);
  text-transform: uppercase;
  margin-bottom: 4px;
}}
.dot {{ display: inline-block; width: 10px; height: 10px; border-radius: 50%; margin-right: 8px; vertical-align: middle; }}

.overlay {{
  display: none; position: fixed; inset: 0;
  background: rgba(15, 23, 42, 0.6); z-index: 3000;
  align-items: center; justify-content: center;
}}
.overlay.v {{ display: flex; }}
.modal {{
  background: var(--white); width: 540px; max-height: 90vh;
  overflow-y: auto; border-radius: 4px;
  box-shadow: 0 20px 50px rgba(0,0,0,0.3);
}}
.mhead {{
  background: var(--clrah-blue);
  padding: 20px 24px; color: var(--white);
  border-bottom: 4px solid var(--clrah-red);
}}
.mtop {{ display: flex; justify-content: space-between; align-items: flex-start; }}
.mname {{ font-family: 'Cinzel', serif; font-size: 1.15rem; font-weight: 700; max-width: 420px; }}
.mpill {{
  display: inline-block; margin-top: 8px;
  font-size: 0.65rem; font-weight: 700; padding: 3px 9px; border-radius: 2px;
  background: rgba(255,255,255,0.2); color: var(--white);
}}
.ppill {{
  display: inline-block; margin-left: 6px;
  font-size: 0.65rem; font-weight: 700; padding: 3px 9px; border-radius: 2px;
  background: var(--clrah-red); color: var(--white);
}}
.xbtn {{
  background: rgba(255,255,255,0.2); border: none; color: var(--white);
  font-size: 0.9rem; width: 28px; height: 28px; border-radius: 3px; cursor: pointer;
}}
.campo {{
  display: grid; grid-template-columns: 120px 1fr; gap: 10px;
  padding: 12px 24px; border-bottom: 1px solid var(--rule); font-size: 0.82rem;
}}
.campo:last-child {{ border-bottom: none; }}
.clabel {{ font-size: 0.65rem; font-weight: 700; color: var(--clrah-gray); text-transform: uppercase; }}
.cval {{ color: var(--ink); line-height: 1.5; }}
.macc {{ padding: 16px 24px; display: flex; gap: 10px; flex-wrap: wrap; border-top: 1px solid var(--rule); }}
.btn {{
  font-family: 'Roboto', sans-serif; font-size: 0.75rem; font-weight: 700;
  padding: 8px 16px; border: none; border-radius: 3px; cursor: pointer; text-decoration: none;
}}
.bp {{ background: var(--clrah-blue); color: var(--white); }}
.bs {{ background: #15803D; color: var(--white); }}
.bd {{ background: var(--clrah-red); color: var(--white); }}
.bg {{ background: var(--rule); color: var(--ink); margin-left: auto; }}

.apanel {{
  display: none; position: fixed; inset: 0;
  background: rgba(15, 23, 42, 0.6); z-index: 3001;
  align-items: center; justify-content: center;
}}
.apanel.v {{ display: flex; }}
.abox {{ background: var(--white); width: 600px; max-height: 90vh; overflow-y: auto; border-radius: 4px; }}
.ahead {{
  background: var(--clrah-blue); color: var(--white);
  padding: 18px 24px; display: flex; justify-content: space-between; align-items: center;
  border-bottom: 4px solid var(--clrah-red);
}}
.ahead h3 {{ font-family: 'Cinzel', serif; font-size: 1rem; }}
.abody {{ padding: 22px; }}
.abody h4 {{ font-family: 'Cinzel', serif; font-size: 0.85rem; color: var(--clrah-blue); margin: 16px 0 6px; }}
.abody p {{ font-size: 0.8rem; color: var(--muted); margin-bottom: 8px; line-height: 1.6; }}
.abody pre {{ background: #F1F5F9; border-left: 3px solid var(--clrah-blue); padding: 12px; font-size: 0.73rem; }}
code {{ color: var(--clrah-red); font-weight: 700; }}
</style>
</head>
<body>

<header>
  <div class="h-izq">
    <span class="h-marca">c<span>L</span>RAH</span>
    <span class="h-titulo">Centro Logístico Regional de Asistencia Humanitaria — Vitral de Proveedores</span>
  </div>
  <div class="h-der">
    <span class="h-badge" id="hbadge">{n} proveedores</span>
    <button class="h-btn" onclick="document.getElementById('ap').classList.add('v')">+ Agregar proveedor</button>
  </div>
</header>

<div class="filtros">
  <span class="flabel">Categoría</span>
  <select id="fc" onchange="filtrar()">
    <option value="">Todas</option>
    <option value="wash">WASH (Agua, Saneamiento e Higiene)</option>
    <option value="shelter">Emergency Shelter (Refugio Temporal)</option>
    <option value="camp">Camp Management (Gestión Campamentos)</option>
    <option value="protection">Protection (Protección e Inclusión)</option>
    <option value="health">Health (Salud y Farmacéutico)</option>
    <option value="logistica">Servicio de Transporte y Logística</option>
    <option value="otro">Otro / Insumos General</option>
  </select>
  <div class="fsep"></div>
  <span class="flabel">Transporte</span>
  <select id="ft" onchange="filtrar()">
    <option value="">Todos</option>
    <option value="Multimodal">Multimodal</option>
    <option value="Terrestre">Terrestre</option>
    <option value="aéreo">Incluye aéreo</option>
  </select>
  <div class="fsep"></div>
  <input type="text" id="fb" placeholder="Buscar proveedor..." oninput="filtrar()">
  <span id="rcount"><span id="nv">{n}</span> de {n}</span>
</div>

<div class="grid">
  <div class="sidebar" id="sb"></div>
  <div id="mapa"></div>
</div>

<!-- Modal detalle -->
<div class="overlay" id="ov" onclick="cerrarSi(event)">
  <div class="modal">
    <div class="mhead">
      <div class="mtop">
        <div class="mname" id="mn"></div>
        <button class="xbtn" onclick="cerrar()">✕</button>
      </div>
      <div style="margin-top:8px;">
        <span class="mpill" id="mp"></span>
        <span class="ppill" id="mpr" style="display:none">Prioridad Alta</span>
      </div>
    </div>
    <div>
      <div class="campo"><span class="clabel">Región</span><span class="cval" id="mr"></span></div>
      <div class="campo"><span class="clabel">Transporte</span><span class="cval" id="mt"></span></div>
      <div class="campo"><span class="clabel">Capacidad</span><span class="cval" id="mc"></span></div>
      <div class="campo"><span class="clabel">Certificaciones</span><span class="cval" id="mce"></span></div>
      <div class="campo"><span class="clabel">Estado</span><span class="cval" id="me"></span></div>
      <div class="campo"><span class="clabel">Teléfono</span><span class="cval" id="mte"></span></div>
      <div class="campo"><span class="clabel">Correo</span><span class="cval" id="mem"></span></div>
      <div class="campo"><span class="clabel">Notas</span><span class="cval" id="mno"></span></div>
    </div>
    <div class="macc" id="ma"></div>
  </div>
</div>

<!-- Panel agregar proveedor -->
<div class="apanel" id="ap" onclick="cerrarAp(event)">
  <div class="abox">
    <div class="ahead">
      <h3>Cómo agregar o actualizar proveedores</h3>
      <button class="xbtn" onclick="document.getElementById('ap').classList.remove('v')">✕</button>
    </div>
    <div class="abody">
      <h4>Este vitral se nutre de la base de datos Excel</h4>
      <p>El master es el archivo <code>Base_de_datos_Vitral_de_Proveedores.xlsx</code>.</p>

      <h4>Para agregar o editar una categoría</h4>
      <p>Usa la lista desplegable en la columna <code>Categoría de Servicio</code> eligiendo entre las opciones estándar (WASH, Emergency Shelter, Camp Management, Protection, Health, Servicio de Transporte y Logística, u Otro).</p>

      <h4>Columnas geográficas requeridas</h4>
      <pre>Columna M → Latitud       (ej: 9.3548)
Columna N → Longitud      (ej: -79.8994)
Columna O → Lugar (mapa)  (ej: France Field, ZLC)
Columna P → Prioridad     (SI o NO)</pre>

      <div style="margin-top:20px;">
        <button class="btn bp" onclick="document.getElementById('ap').classList.remove('v')">Entendido</button>
      </div>
    </div>
  </div>
</div>

<script>
const P = {datos_js};

const map = L.map('mapa').setView([9.15, -79.75], 9);

L.tileLayer('https://server.arcgisonline.com/ArcGIS/rest/services/World_Street_Map/MapServer/tile/{{z}}/{{y}}/{{x}}', {{
  attribution: 'Tiles &copy; Esri &mdash; CLRAH Vitral',
  maxZoom: 18
}}).addTo(map);

setTimeout(() => {{ map.invalidateSize(); }}, 300);

const cols = {{
  wash: '#0284C7',
  shelter: '#16A34A',
  camp: '#CA8A04',
  protection: '#DC2626',
  health: '#7C3AED',
  logistica: '#0055A5',
  otro: '#64748B'
}};

function icono(p) {{
  const c = cols[p.catKey] || '#64748B', sz = p.prioridad ? 16 : 10;
  return L.divIcon({{
    className: '',
    html: `<div style="width:${{sz}}px;height:${{sz}}px;border-radius:50%;background:${{c}};border:2.5px solid white;box-shadow:0 1px 6px rgba(0,0,0,.45);${{p.prioridad ? 'outline:2.5px solid #E30613;outline-offset:2px;' : ''}}"></div>`,
    iconSize: [sz, sz], iconAnchor: [sz/2, sz/2]
  }});
}}

const ley = L.control({{ position: 'bottomright' }});
ley.onAdd = () => {{
  const d = L.DomUtil.create('div', 'lmapa');
  d.innerHTML = `<b>Categoría de Servicio</b>
    <div><span class="dot" style="background:#0284C7"></span>WASH</div>
    <div><span class="dot" style="background:#16A34A"></span>Emergency Shelter</div>
    <div><span class="dot" style="background:#CA8A04"></span>Camp Management</div>
    <div><span class="dot" style="background:#DC2626"></span>Protection</div>
    <div><span class="dot" style="background:#7C3AED"></span>Health</div>
    <div><span class="dot" style="background:#0055A5"></span>Transporte y Logística</div>
    <div><span class="dot" style="background:#64748B"></span>Otro / Insumos General</div>
    <div style="margin-top:6px;font-size:.66rem;color:#E30613;font-weight:700">Borde rojo = Prioridad Alta</div>`;
  return d;
}};
ley.addTo(map);

let marks = [], activo = null;

function renderSb(lista) {{
  const sb = document.getElementById('sb'); sb.innerHTML = '';
  lista.forEach(p => {{
    const d = document.createElement('div');
    d.className = 'item' + (activo === p ? ' activo' : '');
    d.innerHTML = `${{p.prioridad ? '<span class="pmark">PRIORIDAD</span>' : ''}}<div class="iname">${{p.nombre}}</div><div class="imeta"><span class="cpill ${{p.catKey}}">${{p.cat}}</span><span class="ilugar">${{p.lugar}}</span></div>`;
    d.onclick = () => abrir(p, d);
    sb.appendChild(d);
  }});
  document.getElementById('nv').textContent = lista.length;
  document.getElementById('hbadge').textContent = lista.length + ' proveedores';
}}

function renderMap(lista) {{
  marks.forEach(m => map.removeLayer(m)); marks = [];
  lista.forEach(p => {{
    const m = L.marker([p.lat, p.lng], {{ icon: icono(p) }})
      .addTo(map)
      .bindPopup(`<b style="font-family:'Cinzel',serif;font-size:.85rem;color:#0F172A">${{p.nombre}}</b><br><span style="font-size:.72rem;color:#64748B">${{p.lugar}}</span>${{p.prioridad ? '<br><b style="color:#E30613;font-size:.7rem">★ Prioridad Alta</b>' : ''}}`, {{ maxWidth: 220 }})
      .on('click', () => abrir(p));
    marks.push(m);
  }});
}}

function filtrar() {{
  const fc = document.getElementById('fc').value,
        ft = document.getElementById('ft').value,
        fb = document.getElementById('fb').value.toLowerCase();
  const l = P.filter(p =>
    (!fc || p.catKey === fc) &&
    (!ft || p.trans.toLowerCase().includes(ft.toLowerCase())) &&
    (!fb || p.nombre.toLowerCase().includes(fb))
  );
  renderSb(l); renderMap(l);
}}

function abrir(p, el) {{
  activo = p;
  document.querySelectorAll('.item').forEach(x => x.classList.remove('activo'));
  if (el) el.classList.add('activo');
  document.getElementById('mn').textContent = p.nombre;
  document.getElementById('mp').textContent = p.cat;
  document.getElementById('mr').textContent = p.region || '—';
  document.getElementById('mt').textContent = p.trans || '—';
  document.getElementById('mc').textContent = p.cap || '—';
  document.getElementById('mce').textContent = p.cert || '—';
  document.getElementById('me').textContent = p.estado || '—';
  document.getElementById('mte').textContent = p.tel || '—';
  document.getElementById('mem').textContent = p.email || '—';
  document.getElementById('mno').textContent = p.notas || '—';
  document.getElementById('mpr').style.display = p.prioridad ? 'inline-block' : 'none';
  const ac = document.getElementById('ma'); ac.innerHTML = '';
  const ems = (p.email || '').split('/').map(e => e.trim()).filter(e => e.includes('@'));
  if (ems.length) {{ const a = document.createElement('a'); a.href = `mailto:${{ems[0]}}`; a.className = 'btn bp'; a.innerHTML = '✉ Enviar correo'; ac.appendChild(a); }}
  const tl = (p.tel || '').replace(/[^0-9+]/g, '');
  if (tl.length >= 7) {{ const wa = document.createElement('a'); wa.href = `https://wa.me/${{tl.startsWith('+') ? tl.slice(1) : tl}}`; wa.target = '_blank'; wa.className = 'btn bs'; wa.innerHTML = ' WhatsApp'; ac.appendChild(wa); }}
  const gm = document.createElement('a'); gm.href = `https://maps.google.com/?q=${{encodeURIComponent(p.nombre + ' ' + p.lugar + ' Panamá')}}`; gm.target = '_blank'; gm.className = 'btn bd'; gm.innerHTML = ' Ver en Maps'; ac.appendChild(gm);
  const cx = document.createElement('button'); cx.className = 'btn bg'; cx.textContent = 'Cerrar'; cx.onclick = cerrar; ac.appendChild(cx);
  document.getElementById('ov').classList.add('v');
  map.setView([p.lat, p.lng], 13, {{ animate: true }});
}}

function cerrar() {{ document.getElementById('ov').classList.remove('v'); }}
function cerrarSi(e) {{ if (e.target === document.getElementById('ov')) cerrar(); }}
function cerrarAp(e) {{ if (e.target === document.getElementById('ap')) document.getElementById('ap').classList.remove('v'); }}

renderSb(P);
renderMap(P);
</script>
</body>
</html>"""
    return html

# ── Main ──────────────────────────────────────────────────────
if __name__ == "__main__":
    if not EXCEL_FILE.exists():
        print(f"Error: no se encontró '{EXCEL_FILE.name}'")
        sys.exit(1)

    print(f"Leyendo: {EXCEL_FILE.name}")
    proveedores = leer_excel()

    print(f"Generando: {HTML_FILE.name}")
    html = generar_html(proveedores)
    HTML_FILE.write_text(html, encoding="utf-8")

    print(f"\n✓ Listo — {HTML_FILE.name} actualizado con {len(proveedores)} proveedores")